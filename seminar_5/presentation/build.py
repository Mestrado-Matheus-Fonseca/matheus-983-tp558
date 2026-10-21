"""Build the editable deck, speaker notes and PDF with the previous seminar identity."""

import json
from pathlib import Path
import subprocess

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
INK, BLUE, MUTED, PALE = "172033", "004F8B", "64748B", "EFF6FF"


def text(slide, content, x, y, width, height, size=20, color=INK, bold=False):
    """Add a consistent text box.

    Args:
        slide: Target slide.
        content: Text content.
        x: Left coordinate in inches.
        y: Top coordinate in inches.
        width: Width in inches.
        height: Height in inches.
        size: Font size in points.
        color: RGB hex color.
        bold: Whether to use bold type.

    Returns:
        The created textbox shape.
    """
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    for index, line in enumerate(content.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Verdana"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(8)
    return shape


def box(slide, title, body, x, y, width, height=3.6):
    """Add a restrained information card matching earlier seminar colors.

    Args:
        slide: Target slide.
        title: Short card heading.
        body: Card content.
        x: Left coordinate.
        y: Top coordinate.
        width: Card width.
        height: Card height.
    """
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(PALE)
    shape.line.color.rgb = RGBColor.from_string("D5DEE8")
    text(slide, title, x + .22, y + .23, width - .44, .75, 21, BLUE, True)
    text(slide, body, x + .22, y + 1.15, width - .44, height - 1.35, 20)


def picture(slide, path, x=.7, y=1.65, width=11.9, height=4.7):
    """Fit an image without changing its aspect ratio.

    Args:
        slide: Target slide.
        path: Local image path.
        x: Left bound.
        y: Top bound.
        width: Available width.
        height: Available height.
    """
    with Image.open(path) as image:
        ratio = image.width / image.height
    actual_width = min(width, height * ratio)
    actual_height = actual_width / ratio
    slide.shapes.add_picture(str(path), Inches(x + (width - actual_width) / 2),
                             Inches(y + (height - actual_height) / 2),
                             width=Inches(actual_width), height=Inches(actual_height))


def main():
    """Generate a 16:9 PPTX with embedded videos and export its static PDF."""
    specs = json.loads((ROOT / "presentation/slides.json").read_text())
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    notes = ["# Roteiro do Seminário 5", "",
             "Duração sugerida: 30–35 minutos. Vídeos incorporados por clique.",
             "Texto detalhado nas notas do PPTX e nos documentos científicos.", ""]
    for index, item in enumerate(specs, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        text(slide, item["section"], .1, .08, 12.8, .3, 11, BLUE)
        text(slide, item["title"], .5, .58, 12.3, .85, 29, INK, True)
        text(slide, "Inatel • TP558", .5, 7.02, 3, .3, 10, BLUE, True)
        number = text(slide, str(index), 12.3, 7.02, .5, .3, 10, MUTED)
        number.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        source = item.get("source", "")
        text(slide, source, .7, 6.76, 11.8, .25, 9, MUTED)
        kind = item["kind"]
        if kind == "cover":
            text(slide, "TP558 · Tópicos avançados em Machine Learning",
                 .7, 1.9, 11.8, .6, 23, BLUE)
            text(slide, item["subtitle"], .7, 2.8, 11.8, 1.3, 30)
            text(slide, "Matheus Henrique Fonseca Afonso\nmatheus.fonseca@mtel.inatel.br",
                 .7, 5.2, 11.8, 1.0, 19, MUTED)
        elif kind in {"cards", "references", "equation"}:
            y, height = (3.0, 3.3) if kind == "equation" else (1.85, 4.2)
            if kind == "equation":
                text(slide, item["formula"], .8, 1.7, 11.7, .8, 34, BLUE, True)
            for col, (title, body) in enumerate(item["cards"]):
                box(slide, title, body, .65 + col * 4.08, y, 3.88, height)
        elif kind == "figure":
            picture(slide, ROOT / item["figure"])
            text(slide, item.get("takeaway", ""), .7, 6.34, 11.9, .4, 17, MUTED)
        elif kind == "paper":
            picture(slide, ROOT / item["figure"], x=.7, y=1.5, width=5.1, height=5.05)
            box(slide, "Evidência do artigo", item["takeaway"], 6.5, 1.9, 5.8, 2.2)
            text(slide, "Página original disponível em maior resolução nos assets.\n"
                 "Curvas sem tabela numérica não foram transcritas.",
                 6.7, 4.6, 5.3, 1.4, 20, MUTED)
        elif kind == "video":
            scene = item["scene"]
            slide.shapes.add_movie(
                str(ROOT / f"assets/videos/{scene}.mp4"),
                Inches(2.02), Inches(1.4), Inches(9.3), Inches(5.23),
                poster_frame_image=str(ROOT / f"assets/images/{scene}.png"),
                mime_type="video/mp4",
            )
        elif kind == "comparison":
            rows = [
                ["Medida", "Artigo", "Demonstração", "SAR"],
                ["FID, 512/4 passos", "11,10 (LCM)", "N/A", "Não avaliado"],
                ["CLIP, 512/4 passos", "28,69 (LCM)", "N/A", "Não avaliado"],
                ["Aderência à classe", "Não reportada", "CSV local", "Não aplicável"],
                ["IoU / Dice", "Não reportados", "Não avaliados", "Experimento futuro"],
            ]
            table = slide.shapes.add_table(len(rows), 4, Inches(.7), Inches(1.8),
                                           Inches(11.9), Inches(3.7)).table
            for row_id, row in enumerate(rows):
                for col_id, value in enumerate(row):
                    cell = table.cell(row_id, col_id)
                    cell.text = value
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(BLUE if row_id == 0 else PALE)
                    for p in cell.text_frame.paragraphs:
                        p.font.name = "Verdana"
                        p.font.size = Pt(18)
                        p.font.color.rgb = RGBColor.from_string("FFFFFF" if row_id == 0 else INK)
            text(slide, "N/A: configuração experimental diferente", .7, 5.85, 11.8, .6, 23, MUTED)
        elif kind == "research_pipeline":
            labels = [("Treino real", "SAR + máscaras"),
                      ("Professor", "DDPM condicionado"),
                      ("Aluno", "Consistência / LCM")]
            for col, (title, body) in enumerate(labels):
                box(slide, title, body, .65 + col * 4.08, 1.65, 3.88, 2.3)
                if col < 2:
                    text(slide, "→", 4.43 + col * 4.08, 2.4, .4, .5, 24, BLUE)
            box(slide, "Pares sintéticos auditados", "Imagem e máscara alinhadas. Apenas treino.",
                .65, 4.4, 5.85, 2.1)
            box(slide, "Mesmo SAM / YOLO", "Avaliação em SAR real intocado: IoU e Dice.",
                6.8, 4.4, 5.85, 2.1)
        else:
            raise ValueError(f"Tipo de slide desconhecido: {kind}")
        slide.notes_slide.notes_text_frame.text = item["notes"] + "\n\nFonte: " + source
        notes.extend([f"## {index:02d}. {item['title']}", "", item["notes"], "",
                      f"Fonte: {source}" if source else "", ""])
    target = ROOT / "presentation/TP558_LCM.pptx"
    presentation.save(target)
    (ROOT / "presentation/speaker_notes.md").write_text("\n".join(notes))
    profile = ROOT / ".cache/libreoffice"
    profile.mkdir(parents=True, exist_ok=True)
    subprocess.run(["libreoffice", f"-env:UserInstallation={profile.as_uri()}", "--headless",
                    "--convert-to", "pdf", "--outdir", str(target.parent), str(target)],
                   check=True, timeout=120)
    if not target.with_suffix(".pdf").exists():
        raise RuntimeError("LibreOffice não produziu o PDF")
    print(f"{len(specs)} slides, notas, cinco vídeos incorporados e PDF gerados.", flush=True)


if __name__ == "__main__":
    main()
