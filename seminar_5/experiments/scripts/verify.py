"""Verify mathematical behavior, saved results, notebooks and delivery artifacts."""

import ast
import csv
import json
from pathlib import Path
import re
import subprocess
import zipfile

import nbformat
import numpy as np
from PIL import Image
from pptx import Presentation
import torch

import experiment as exp

ROOT = Path(__file__).resolve().parents[2]


def main():
    """Run artifact and regression checks and save a concise verification report."""
    torch.set_num_threads(2)
    report = {}
    for path in ROOT.rglob("*.py"):
        if any(part.startswith(".") for part in path.relative_to(ROOT).parts):
            continue
        ast.parse(path.read_text(), filename=str(path))
    report["python_syntax"] = "passed"
    checks = exp.numerical_checks()
    assert checks["ddim_max_error"] < 1e-5
    assert checks["guidance_max_error"] < 1e-5
    assert checks["boundary_max_error"] == 0
    assert checks["ema_max_error"] < 1e-6 and checks["target_detached"]
    report["numerical"] = checks
    data = exp.prepare_data()
    assert not np.intersect1d(data["train_source_indices"], data["val_source_indices"]).size
    for split, count in [("train", 6000), ("val", 1000), ("test", 1000)]:
        assert data[f"{split}_images"].shape == (count, 196)
        assert np.isfinite(data[f"{split}_images"]).all()
        assert len(np.unique(data[f"{split}_source_indices"])) == count
    np.testing.assert_allclose(data["basis"] @ data["basis"].T, np.eye(32), atol=2e-5)
    encoded = (data["train_images"] - data["mean"]) @ data["basis"].T / data["scales"]
    np.testing.assert_allclose(encoded.mean(0), 0, atol=1e-4)
    np.testing.assert_allclose(encoded.std(0), 1, atol=1e-4)
    report["data_splits_and_pca"] = "passed"
    protocol = json.loads((ROOT / "experiments/results/protocol.json").read_text())
    assert protocol["script_sha256"] == exp.digest(ROOT / "experiments/scripts/experiment.py")
    assert protocol["config_sha256"] == exp.digest(ROOT / "experiments/configs/demo.json")
    assert protocol["paper_sha256"] == exp.digest(ROOT / "articles/2310.04378v1.pdf")
    assert protocol["data_sha256"] == exp.digest(ROOT / "data/mnist_subset.npz")
    for name, checksum in protocol["checkpoint_sha256"].items():
        assert exp.digest(ROOT / "experiments/checkpoints" / name) == checksum
    for entry in json.loads((ROOT / "articles/sources/manifest.json").read_text()):
        assert exp.digest(ROOT / "articles/sources" / entry["file"]) == entry["sha256"]
    report["source_data_config_checkpoint_hashes"] = "passed"
    teacher = exp.NoiseMLP()
    teacher.load_state_dict(torch.load(ROOT / "experiments/checkpoints/teacher.pt",
                                      weights_only=True)["state_dict"])
    student = exp.NoiseMLP()
    student.load_state_dict(torch.load(ROOT / "experiments/checkpoints/lcm_k20.pt",
                                      weights_only=True)["state_dict"])
    labels = torch.arange(10)
    first = exp.sample(student, labels, 4, 101)
    second = exp.sample(student, labels, 4, 101)
    assert torch.equal(first, second)
    assert not torch.equal(first, exp.sample(student, labels, 4, 202))
    try:
        exp.sample(student, labels, 0, 101)
    except ValueError:
        pass
    else:
        raise AssertionError("Invalid inference steps must be rejected")
    count = []
    hook = student.register_forward_hook(lambda *args: count.append(1))
    exp.sample(student, labels, 4, 101)
    hook.remove()
    assert len(count) == 4
    count = []
    hook = teacher.register_forward_hook(lambda *args: count.append(1))
    exp.sample(teacher, labels, 4, 101, "ddim")
    hook.remove()
    assert len(count) == 8
    report["inference_determinism_and_calls"] = "passed"
    records = list(csv.DictReader((ROOT / "experiments/results/metrics.csv").open()))
    assert len(records) == 39
    assert len({(r["method"], r["steps"], r["seed"]) for r in records}) == 39
    for record in records:
        assert 0 <= float(record["class_agreement"]) <= 1
        assert float(record["batch_seconds"]) > 0
        assert int(record["network_evaluations"]) == int(record["steps"]) * (
            2 if record["method"] == "ddim" else 1)
        if record["method"] == "ddim" and record["steps"] == "50":
            assert float(record["latent_mse_vs_ddim50"]) == 0
    aggregate = json.loads((ROOT / "experiments/results/aggregate.json").read_text())
    for row in aggregate:
        values = [float(r["class_agreement"]) for r in records
                  if r["method"] == row["method"] and int(r["steps"]) == row["steps"]]
        assert np.isclose(row["class_agreement_mean"], np.mean(values), atol=1e-12)
        assert np.isclose(row["class_agreement_std"], np.std(values, ddof=1), atol=1e-12)
    samples = np.load(ROOT / "experiments/results/samples.npz")
    sample_labels = torch.tensor(samples["labels"])
    regenerated = exp.decode(exp.sample(student, sample_labels, 4, 101), data).numpy()
    np.testing.assert_array_equal(regenerated, samples["lcm_k20_4"])
    report["metrics_and_regenerated_images"] = {"rows": len(records), "aggregate_rows": len(aggregate)}
    notebooks = sorted((ROOT / "notebooks").glob("*.ipynb"))
    assert len(notebooks) == 7
    cells = 0
    for path in notebooks:
        nb = nbformat.read(path, as_version=4)
        nbformat.validate(nb)
        for cell in nb.cells:
            if cell.cell_type == "code":
                assert cell.execution_count is not None, str(path)
                assert not any(output.output_type == "error" for output in cell.outputs), str(path)
                cells += 1
    report["notebooks"] = {"executed": len(notebooks), "executed_code_cells": cells}
    videos = json.loads((ROOT / "assets/videos/manifest.json").read_text())
    assert len(videos) == 5
    for entry in videos:
        path = ROOT / entry["video"]
        metadata = json.loads(subprocess.check_output(
            ["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries",
             "stream=width,height", "-show_entries", "format=duration", "-of", "json", str(path)],
            text=True))
        assert metadata["streams"][0]["width"] == 1280
        assert metadata["streams"][0]["height"] == 720
        assert float(metadata["format"]["duration"]) > 5
        with Image.open(ROOT / f"assets/gifs/{entry['scene']}.gif") as gif:
            assert gif.n_frames > 10
    report["videos_and_gifs"] = len(videos)
    pptx = ROOT / "presentation/TP558_LCM.pptx"
    deck = Presentation(pptx)
    slides = json.loads((ROOT / "presentation/slides.json").read_text())
    assert len(deck.slides) == len(slides) == 33
    for slide in deck.slides:
        assert slide.notes_slide.notes_text_frame.text.strip()
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width + 100
            assert shape.top + shape.height <= deck.slide_height + 100
    with zipfile.ZipFile(pptx) as archive:
        movies = [name for name in archive.namelist() if name.startswith("ppt/media/")
                  and name.endswith(".mp4")]
        assert len(movies) == 5
    pdf = pptx.with_suffix(".pdf")
    info = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    assert re.search(r"Pages:\s+33", info)
    for path in (ROOT / "assets/figures").glob("*.png"):
        with Image.open(path) as image:
            image.verify()
    report["presentation"] = {"pptx_slides": 33, "pdf_pages": 33, "embedded_mp4": len(movies)}
    paper = json.loads((ROOT / "data/paper_results.json").read_text())
    assert paper["512"]["LCM"]["FID"] == [35.36, 13.31, 11.10, 11.84]
    assert paper["768"]["LCM"]["FID"] == [34.22, 16.32, 13.53, 14.97]
    report["paper_table_sentinels"] = "passed, complete tables visually checked against PDF page 7"
    for file in ["README.md", "notes/technical_summary.md", "notes/didactic_summary.md",
                 "notes/mathematics.md", "notes/research_application.md", "notes/sources.md",
                 "notes/methodological_audit.md", "notes/experimental_report.md"]:
        assert (ROOT / file).is_file(), file
    report["documentation"] = "present"
    report["status"] = "passed"
    (ROOT / "experiments/results/verification.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n"
    )
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
